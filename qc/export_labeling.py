"""Export a stratified, blind labeling sample for domain experts. Read-only.

Sheet "Label" shows only record content (no scores) so labels are not biased by our decision.
Sheet "Key" holds our scores and decision, keyed by row_id, for later calibration.

Usage: python -m qc.export_labeling [--n 60] [--seed 7] [--out out/labeling_sample.xlsx]
"""
import argparse
import csv
import random
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from qc.db import connect, get_schema

# (name, how many, predicate on a results row)
STRATA = [
    ("off_topic_flag",    14, lambda r: "possibly_off_topic" in r["major"]),
    ("weak_topic_flag",   10, lambda r: "weak_topic_match" in r["flags"] and "possibly_off_topic" not in r["major"]),
    ("short_abstract",     6, lambda r: "abstract_very_short" in r["major"] and "possibly_off_topic" not in r["major"]),
    ("review_borderline", 10, lambda r: r["decision"] == "review" and not r["major"]),
    ("publish_high",       8, lambda r: r["decision"] == "publish" and float(r["overall"]) >= 0.90),
    ("publish_low",        8, lambda r: r["decision"] == "publish" and float(r["overall"]) < 0.80),
    ("reject",             4, lambda r: r["decision"] == "reject"),
]


def pick_spread(rows, k, rng, taken):
    """Pick k rows round-robin across sources so one big source does not dominate."""
    by_src = defaultdict(list)
    for r in rows:
        if r["id"] not in taken:
            by_src[r["source"]].append(r)
    for v in by_src.values():
        rng.shuffle(v)
    srcs = list(by_src); rng.shuffle(srcs)
    out = []
    while len(out) < k and any(by_src.values()):
        for s in srcs:
            if by_src[s] and len(out) < k:
                out.append(by_src[s].pop())
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=60)
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--out", default="out/labeling_sample.xlsx")
    args = ap.parse_args()
    rng = random.Random(args.seed)

    with open("out/qc_results.csv", encoding="utf-8", newline="") as f:
        results = list(csv.DictReader(f))

    scale = args.n / sum(k for _, k, _ in STRATA)
    chosen, taken = [], set()
    for name, k, pred in STRATA:
        for r in pick_spread([r for r in results if pred(r)], max(1, round(k * scale)), rng, taken):
            taken.add(r["id"]); chosen.append((name, r))
    rng.shuffle(chosen)

    s = get_schema()
    ids = [r["id"] for _, r in chosen]
    with connect() as conn, conn.cursor() as cur:
        cur.execute("SET statement_timeout = '300s'")
        cur.execute(f"SELECT id::text, abstract, url, language FROM {s}.records WHERE id = ANY(%s::uuid[])", (ids,))
        extra = {i: (a, u, l) for i, a, u, l in cur.fetchall()}

    wb = Workbook()
    ins = wb.active; ins.title = "Instructions"
    for line in [
        "Labeling sample for the SoilWise Knowledge Hub quality check",
        "",
        f"{len(chosen)} records, shuffled. Scores and our automated decision are hidden on purpose (see sheet 'Key' only after labeling).",
        "For each row in 'Label', fill the three yellow columns:",
        "  relevant:   yes = clearly about soil / land / agriculture; partly = related but marginal; no = unrelated.",
        "  publishable: yes = metadata good enough to show in the Hub; fixes = relevant but metadata needs work; no = should not be published.",
        "  comment: optional (why, or what is wrong).",
        "Judge the title and abstract shown; open the URL if unsure. Work alone, do not discuss with others (so labels stay independent).",
        "Return the file to the project team when done.",
    ]:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=13); ins.column_dimensions["A"].width = 140

    lab = wb.create_sheet("Label")
    head = ["row_id", "title", "abstract", "source", "resource_type", "language", "url",
            "relevant", "publishable", "comment"]
    lab.append(head)
    key = wb.create_sheet("Key")
    key.append(["row_id", "record_id", "stratum", "decision", "overall", "completeness", "consistency",
                "credibility", "relevance", "accuracy", "flags", "major"])

    for n, (stratum, r) in enumerate(chosen, 1):
        abstract, url, lang = extra.get(r["id"], ("", "", ""))
        lab.append([n, r["title"], (abstract or "")[:1500], r["source"], r["resource_type"], lang, url or "", "", "", ""])
        key.append([n, r["id"], stratum, r["decision"], r["overall"], r["completeness"], r["consistency"],
                    r["credibility"], r["relevance"], r["accuracy"], r["flags"], r["major"]])

    for c, w in zip("ABCDEFGHIJ", [8, 45, 90, 14, 16, 10, 40, 12, 14, 40]):
        lab.column_dimensions[c].width = w
    for cell in lab[1]:
        cell.font = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for row in lab.iter_rows(min_row=2, max_row=lab.max_row):
        for cell in row[:7]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row[7:]:
            cell.fill = yellow
    for cell in lab[1][7:]:
        cell.fill = yellow
    dv1 = DataValidation(type="list", formula1='"yes,partly,no"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"yes,fixes,no"', allow_blank=True)
    lab.add_data_validation(dv1); lab.add_data_validation(dv2)
    dv1.add(f"H2:H{lab.max_row}"); dv2.add(f"I2:I{lab.max_row}")
    lab.freeze_panes = "C2"
    for cell in key[1]:
        cell.font = Font(bold=True)

    wb.save(args.out)
    from collections import Counter
    print(f"Wrote {args.out}: {len(chosen)} records")
    print(dict(Counter(st for st, _ in chosen)))
    print(dict(Counter(r['source'] for _, r in chosen)))


if __name__ == "__main__":
    main()
